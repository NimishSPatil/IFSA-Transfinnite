"""
export_excel.py  -  bundles the master tables into one workbook judges can open.

    python code/export_excel.py   ->   PulpRiot_Master.xlsx

Sheets: README, Case_Facts, Consistency_Checks, Gaps, Assumptions, nodes, edges, products, bom,
channels, events, loss_ledger, risk_flags, demand_monthly (pivot of volume by month x flavor x
depot x channel), inventory_snapshot.  The 16,470-row daily sales table stays in the CSV/SQLite
(demand_daily) to keep the workbook light.
"""
from pathlib import Path

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent.parent
DATA = ROOT / "data"
read = lambda n: pd.read_csv(DATA / f"{n}.csv", encoding="utf-8-sig")

demand = read("demand_daily")
demand["month"] = demand["date"].str[:7]
monthly = demand.groupby(["month", "flavor_id", "depot_id", "channel_id"], as_index=False)["units"].sum()

readme = pd.DataFrame({"PulpRiot_Master.xlsx": [
    "Pulp Riot supply network: master data (TransfiNITTe 2026, PS1).",
    "Source of truth is the CSV files in /data; this workbook is generated from them (code/export_excel.py).",
    "data_source column: Given = stated in the case, Derived = calculated from case figures, Assumed = our judgement (see Assumptions).",
    "Every row with a case_fact_id can be traced to Case_Facts; every blank the case leaves is listed in Gaps.",
    "Deep modules: Supply and Inventory. Demand, Distribution and Transportation are covered at a high level.",
    "demand_monthly sums the daily sample sales (9.6 crore bottles). Sample data is built from case figures only; shapes are our assumptions.",
    "risk_flags come from code/build_graph.py. Every recommended action needs human approval (case constraint F71).",
]})

sheets = [("README", readme)] + [(n, read(n)) for n in (
    "Case_Facts", "Consistency_Checks", "Gaps", "Assumptions", "nodes", "edges", "products", "bom",
    "channels", "events", "loss_ledger", "risk_flags")] + [
    ("demand_monthly", monthly), ("inventory_snapshot", read("inventory_snapshot"))]

out = ROOT / "PulpRiot_Master.xlsx"
with pd.ExcelWriter(out, engine="openpyxl") as xw:
    for name, df in sheets:
        df.to_excel(xw, sheet_name=name, index=False)
        ws = xw.sheets[name]
        for c in ws[1]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="1F4E79")
        ws.freeze_panes = "A2"
        for i, col in enumerate(df.columns, 1):
            width = min(max(len(str(col)), *(len(str(v)) for v in df[col].head(200))) + 2, 70)
            ws.column_dimensions[get_column_letter(i)].width = width
            if width >= 60:
                for cell in ws[get_column_letter(i)][1:]:
                    cell.alignment = Alignment(wrap_text=True, vertical="top")
print("wrote", out, "-", len(sheets), "sheets")
